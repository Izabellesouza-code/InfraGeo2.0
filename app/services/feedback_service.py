"""Armazena sugestões enviadas pelo agente InfraGeo."""

from __future__ import annotations

import json
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path

from app.config import get_settings

_lock = threading.Lock()
settings = get_settings()

STATUSES = ("nova", "em_analise", "resolvida")
KINDS = ("sugestao", "reclamacao")
SEED_EMAILS = {
    "marina.costa@exemplo.gov.br",
    "paulo.nogueira@exemplo.gov.br",
    "ana.ribeiro@exemplo.gov.br",
}


def _path() -> Path:
    root = Path(settings.upload_dir).resolve().parent
    root.mkdir(parents=True, exist_ok=True)
    return root / "feedback.json"


EXCEL_HEADERS = [
    "Data",
    "Tipo",
    "Autor",
    "E-mail",
    "Mensagem",
    "Camada",
    "Palavra-chave",
    "Lugar",
    "Latitude",
    "Longitude",
    "BBox",
    "Foto",
    "Status",
    "Id",
]


def _excel_path() -> Path:
    root = Path(settings.upload_dir).resolve().parent
    root.mkdir(parents=True, exist_ok=True)
    return root / "sugestoes_reclamacoes.xlsx"


STATUS_LABELS = {
    "nova": "Nova",
    "em_analise": "Em análise",
    "resolvida": "Resolvida",
}


def _excel_row(item: dict) -> list:
    kind_label = "Reclamação" if item.get("kind") == "reclamacao" else "Sugestão"
    return [
        str(item.get("created_at") or ""),
        kind_label,
        item.get("author") or "",
        item.get("email") or "",
        item.get("message") or "",
        item.get("layer") or "",
        item.get("keyword") or "",
        item.get("place") or "",
        item.get("lat") if item.get("lat") is not None else "",
        item.get("lng") if item.get("lng") is not None else "",
        item.get("bbox") or "",
        item.get("photo_url") or "",
        STATUS_LABELS.get(item.get("status") or "nova", item.get("status") or ""),
        item.get("id") or "",
    ]


def _sync_excel(items: list[dict]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return
    path = _excel_path()
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append(EXCEL_HEADERS)
    rows = [i for i in items if _is_agent_suggestion(i)]
    rows.sort(key=lambda x: x.get("created_at") or "")
    for item in rows:
        ws.append(_excel_row(item))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def ensure_excel() -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return
    path = _excel_path()
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append(EXCEL_HEADERS)
    wb.save(path)


ensure_excel()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _is_agent_suggestion(item: dict) -> bool:
    email = str(item.get("email") or "").lower()
    if email in SEED_EMAILS:
        return False
    if item.get("kind") not in KINDS:
        return False
    source = item.get("source") or "agent"
    return source == "agent"


def _read() -> list[dict]:
    path = _path()
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    return data if isinstance(data, list) else []


def _write(items: list[dict]) -> None:
    path = _path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


def list_items() -> list[dict]:
    with _lock:
        items = [i for i in _read() if _is_agent_suggestion(i)]
    items.sort(key=lambda x: x.get("created_at") or "", reverse=True)
    return items


def add_item(
    *,
    kind: str,
    message: str,
    author: str,
    email: str = "",
    layer: str = "",
    layer_id: str = "",
    keyword: str = "",
    place: str = "",
    lat: float | None = None,
    lng: float | None = None,
    bbox: str = "",
    photo_url: str = "",
) -> dict:
    kind = kind if kind in KINDS else "sugestao"
    text = (message or "").strip()
    if len(text) < 1:
        raise ValueError("Descreva a sugestão.")
    item = {
        "id": str(uuid.uuid4()),
        "kind": kind,
        "status": "nova",
        "source": "agent",
        "author": (author or "Usuário").strip()[:80],
        "email": (email or "").strip()[:120],
        "message": text[:2000],
        "layer": (layer or "").strip()[:200],
        "layer_id": (layer_id or "").strip()[:120],
        "keyword": (keyword or "").strip()[:80],
        "place": (place or "").strip()[:400],
        "lat": lat,
        "lng": lng,
        "bbox": (bbox or "").strip()[:240],
        "photo_url": (photo_url or "").strip()[:500],
        "created_at": _now(),
    }
    with _lock:
        items = _read()
        items.append(item)
        _write(items)
        _sync_excel(items)
    return item


def update_status(item_id: str, status: str) -> dict | None:
    if status not in STATUSES:
        raise ValueError("Status inválido.")
    with _lock:
        items = _read()
        found = None
        for item in items:
            if item.get("id") == item_id:
                item["status"] = status
                item["updated_at"] = _now()
                found = item
                break
        if found:
            _write(items)
            _sync_excel(items)
        return found
